"""材料提纲：把 PDF、Word、PPT、Excel、Markdown 转成短提纲，供 Phase 2 先看提纲再读原文。

用法：python extract_outline.py <文件...> --out 材料提纲/ [--nums 5]

每份材料输出一个 <文件名>.md，写：
  - 章节（PDF 按页，Word 按标题样式，PPT 按页标题，Excel 按工作表，Markdown 按 # 标题）
  - 每节的位置（页码、段号）和字数
  - 每节含数字的句子，最多 --nums 条，每条截到 80 字
  - Word 的表格、Excel 的表头和行列数
读原文时按提纲里的位置只读命中的章节。
依赖：pypdf、python-docx、python-pptx、openpyxl（缺哪个只影响对应格式）。
"""
import argparse
import re
import sys
from pathlib import Path

NUM = re.compile(r"\d[\d,.]*\s*(%|万|亿|元|个|台|人|天|小时|分钟|年|月|次|项|条|吨|件|套|家|倍)?")
SPLIT = re.compile(r"(?<=[。！？；\n])")


def num_lines(text, k):
    out = []
    for s in SPLIT.split(text):
        s = " ".join(s.split())
        if len(s) >= 6 and NUM.search(s) and re.search(r"[一-鿿A-Za-z]", s):
            out.append(s[:80] + ("…" if len(s) > 80 else ""))
            if len(out) >= k:
                break
    return out


def section(title, where, text, k):
    lines = [f"### {title}", f"位置：{where}　字数：{len(''.join(text.split()))}"]
    lines += [f"- {s}" for s in num_lines(text, k)]
    return "\n".join(lines)


def pdf(path, k):
    from pypdf import PdfReader
    r = PdfReader(str(path))
    secs = []
    for i, p in enumerate(r.pages, 1):
        t = p.extract_text() or ""
        first = next((l.strip() for l in t.splitlines() if l.strip()), "（无文字，可能是扫描页）")
        secs.append(section(first[:40], f"p{i}", t, k))
    return f"共 {len(r.pages)} 页", secs


def docx_(path, k):
    import docx
    d = docx.Document(str(path))
    secs, title, start, buf = [], "（开头）", 1, []
    for i, p in enumerate(d.paragraphs, 1):
        st = (p.style.name or "").lower()
        if p.text.strip() and (st.startswith("heading") or st.startswith("标题") or st == "title"):
            if buf or secs:
                secs.append(section(title, f"段 {start}–{i - 1}", "\n".join(buf), k))
            title, start, buf = p.text.strip()[:40], i, []
        else:
            buf.append(p.text)
    secs.append(section(title, f"段 {start}–{len(d.paragraphs)}", "\n".join(buf), k))
    for n, t in enumerate(d.tables, 1):
        head = " | ".join(c.text.strip()[:12] for c in t.rows[0].cells) if t.rows else ""
        secs.append(f"### 表 {n}\n{len(t.rows)} 行 × {len(t.columns)} 列　表头：{head}")
    return f"共 {len(d.paragraphs)} 段、{len(d.tables)} 张表", secs


def pptx_(path, k):
    from pptx import Presentation
    prs = Presentation(str(path))
    secs = []
    for i, s in enumerate(prs.slides, 1):
        texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
        title = s.shapes.title.text_frame.text if s.shapes.title is not None and s.shapes.title.has_text_frame else ""
        if not title:  # 没有标题占位符时取字号最大的文字
            best = (0, "")
            for sh in s.shapes:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    size = max((r.font.size or 0 for p in sh.text_frame.paragraphs for r in p.runs), default=0)
                    best = max(best, (size, sh.text_frame.text.strip()))
            title = best[1] or "（无文字）"
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        secs.append(section(" ".join(title.split())[:40], f"第 {i} 页" + ("，有讲稿" if notes.strip() else ""), "\n".join(texts + [notes]), k))
    return f"共 {len(prs.slides)} 页", secs


def xlsx(path, k):
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    secs = []
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        head = next(rows, None) or ()
        n = 1 + sum(1 for _ in rows)
        cols = ", ".join(str(c)[:12] for c in head if c is not None)
        secs.append(f"### 工作表 {ws.title}\n{n} 行 × {ws.max_column} 列　表头：{cols}")
    return f"共 {len(wb.worksheets)} 个工作表", secs


def md(path, k):
    text = path.read_text(encoding="utf-8", errors="ignore").splitlines()
    secs, title, start, buf = [], "（开头）", 1, []
    for i, l in enumerate(text, 1):
        if l.startswith("#"):
            if buf:
                secs.append(section(title, f"行 {start}–{i - 1}", "\n".join(buf), k))
            title, start, buf = l.lstrip("#").strip()[:40], i, []
        else:
            buf.append(l)
    secs.append(section(title, f"行 {start}–{len(text)}", "\n".join(buf), k))
    return f"共 {len(text)} 行", secs


READERS = {".pdf": pdf, ".docx": docx_, ".pptx": pptx_, ".xlsx": xlsx, ".xlsm": xlsx, ".md": md, ".txt": md}


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("files", nargs="+")
    ap.add_argument("--out", required=True)
    ap.add_argument("--nums", type=int, default=5, help="每节最多列几条含数字的句子")
    a = ap.parse_args()
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    for f in map(Path, a.files):
        reader = READERS.get(f.suffix.lower())
        if not reader:
            print(f"跳过（格式不支持）：{f}", file=sys.stderr)
            continue
        try:
            summary, secs = reader(f, a.nums)
        except Exception as e:  # 单个文件失败不影响其他文件
            print(f"失败：{f}：{e}", file=sys.stderr)
            continue
        dst = out / (f.stem + ".md")
        dst.write_text(f"# {f.name}\n\n{summary}\n\n" + "\n\n".join(secs) + "\n", encoding="utf-8")
        print(f"{dst}　{summary}　原文 {f.stat().st_size // 1024} KB → 提纲 {dst.stat().st_size // 1024} KB")


if __name__ == "__main__":
    main()
